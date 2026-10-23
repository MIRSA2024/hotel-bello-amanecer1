#!/usr/bin/env python3
"""
Inserta las capturas de pantalla en el informe APF1 reemplazando los recuadros
"[Insertar captura de pantalla]" de las Figuras 6, 7, 8, 9, 10, 12, 13 y 14.

Uso:
    pip install python-docx
    python docs/insertar_capturas.py APF_1.docx docs/evidencias APF_1_con_capturas.docx
    # opcional: también pone la URL de tu repositorio
    python docs/insertar_capturas.py APF_1.docx docs/evidencias APF_1_con_capturas.docx \
        --repo-url https://github.com/TU_USUARIO/hotel-bello-amanecer

Las imágenes deben llamarse figura06.png, figura07.png, figura08.png, figura09.png,
figura10.png, figura12.png, figura13.png y figura14.png (también se aceptan .jpg/.jpeg).
Si una figura necesita varias capturas, nómbralas figura07_1.png, figura07_2.png, ...
y se apilan en ese orden. Si falta alguna figura, su recuadro se deja como está y se avisa.
"""
import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm
from docx.table import Table
from docx.text.paragraph import Paragraph

MARCA = "[Insertar captura de pantalla]"
ANCHO_MAX_CM = 15.9  # ancho útil de la página (A4 con márgenes de 2,54 cm)


def buscar_imagenes(carpeta: Path, numero: int):
    """figura07.png, o varias apiladas: figura07_1.png, figura07_2.png, ... (en orden)."""
    patron = re.compile(rf"^figura{numero:02d}(_\d+)?\.(png|jpg|jpeg)$", re.IGNORECASE)
    return sorted((f for f in carpeta.iterdir() if patron.match(f.name)),
                  key=lambda f: [int(x) if x.isdigit() else x for x in re.split(r"(\d+)", f.name.lower())])


def numero_de_figura(elemento, doc):
    """Lee 'Figura N.' del párrafo que sigue a la tabla."""
    siguiente = elemento.getnext()
    if siguiente is None or not siguiente.tag.endswith("}p"):
        return None
    m = re.match(r"\s*Figura\s+(\d+)\.", Paragraph(siguiente, doc).text)
    return int(m.group(1)) if m else None


def reemplazar_url(doc, url):
    cambios = 0
    for p in doc.paragraphs:
        if "<organización>" not in p.text:
            continue
        for run in p.runs:
            nuevo, n = re.subn(r"\[?https://github\.com/<organización>/hotel-bello-amanecer\]?", url, run.text)
            if n:
                run.text = nuevo
                cambios += n
        if cambios == 0:  # texto repartido en varios runs: se reconstruye en el primero
            texto = re.sub(r"\[?https://github\.com/<organización>/hotel-bello-amanecer\]?", url, p.text)
            for r in p.runs[1:]:
                r.text = ""
            p.runs[0].text = texto
            cambios += 1
    return cambios


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("docx_entrada")
    ap.add_argument("carpeta_imagenes")
    ap.add_argument("docx_salida")
    ap.add_argument("--repo-url", help="URL del repositorio de GitHub para reemplazar el marcador del informe")
    a = ap.parse_args()

    doc = Document(a.docx_entrada)
    carpeta = Path(a.carpeta_imagenes)
    puestas, faltan = [], []

    # Se recorre una copia de la lista porque se eliminan elementos mientras se itera
    for el in list(doc.element.body.iterchildren()):
        if not el.tag.endswith("}tbl"):
            continue
        tabla = Table(el, doc)
        if MARCA not in tabla.cell(0, 0).text:
            continue
        n = numero_de_figura(el, doc)
        if n is None:
            continue
        rutas = buscar_imagenes(carpeta, n)
        if not rutas:
            faltan.append(n)
            continue

        # Un párrafo por imagen, justo donde estaba el recuadro (varias imágenes quedan apiladas)
        for ruta in rutas:
            nuevo = doc.add_paragraph()           # se crea al final...
            el.addprevious(nuevo._p)              # ...y se mueve antes del recuadro
            nuevo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            nuevo.paragraph_format.keep_with_next = True  # la imagen no se separa de su título
            nuevo.paragraph_format.space_after = Cm(0.1)
            nuevo.add_run().add_picture(str(ruta), width=Cm(ANCHO_MAX_CM))
        el.getparent().remove(el)
        puestas.append(n)

    if a.repo_url:
        print(f"URL del repositorio reemplazada en {reemplazar_url(doc, a.repo_url)} lugar(es).")

    doc.save(a.docx_salida)
    print("Figuras insertadas:", ", ".join(map(str, puestas)) or "ninguna")
    if faltan:
        print("Faltan imágenes para las figuras:", ", ".join(map(str, faltan)))
    print("Guardado en", a.docx_salida)


if __name__ == "__main__":
    sys.exit(main())
